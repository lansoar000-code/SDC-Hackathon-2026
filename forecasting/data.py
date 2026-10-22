"""读入表格并逐行校验。解析与数据库写入分开，避免导入半张表。"""

import csv
import io
import zipfile
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path
from xml.etree.ElementTree import ParseError

from openpyxl.utils.exceptions import InvalidFileException
from pydantic import ValidationError

from backend.errors import AppError
from backend.schemas import SaleRecord

REQUIRED_COLUMNS = {"date", "dish_id", "dish_name", "sales_quantity"}
OPTIONAL_COLUMNS = {"prepared_quantity", "waste_quantity", "weather", "is_exam_week", "has_event"}


def read_table(content: bytes, filename: str, max_rows: int = 100000) -> list[SaleRecord]:
    """支持 UTF-8 CSV 与第一张工作表 XLSX；返回全部合法记录或整批报错。"""
    suffix = Path(filename).suffix.lower()
    try:
        if suffix == ".csv":
            text = content.decode("utf-8-sig")
            table = csv.reader(io.StringIO(text))
            headers = next(table, [])
            rows = list(_limited_rows(table, max_rows))
        elif suffix == ".xlsx":
            # 限制解压体积，避免很小的压缩文件展开成巨量内容。
            with zipfile.ZipFile(io.BytesIO(content)) as archive:
                if sum(item.file_size for item in archive.infolist()) > 50 * 1024 * 1024:
                    raise AppError("Excel 解压后超过 50 MB，请缩小文件", 413)
            from openpyxl import load_workbook

            workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=False)
            try:
                table = workbook.worksheets[0].iter_rows(values_only=True)
                headers = list(next(table, []))
                rows = list(_limited_rows(table, max_rows))
            finally:
                workbook.close()
        else:
            raise AppError("只支持 .csv 和 .xlsx 文件")
    except AppError:
        raise
    except (
        UnicodeError,
        csv.Error,
        zipfile.BadZipFile,
        ValueError,
        KeyError,
        IndexError,
        EOFError,
        TypeError,
        ParseError,
        InvalidFileException,
    ) as error:
        raise AppError("无法读取表格；CSV 请使用 UTF-8，Excel 请使用有效的 XLSX 文件") from error

    headers = [str(value).strip() if value is not None else "" for value in headers]
    missing = REQUIRED_COLUMNS - set(headers)
    unknown = set(headers) - REQUIRED_COLUMNS - OPTIONAL_COLUMNS
    if missing or unknown or len(headers) != len(set(headers)):
        raise AppError(
            "表头不符合模板", details={"missing": sorted(missing), "unknown": sorted(unknown)}
        )

    records, errors, seen, names = [], [], set(), {}
    for line, values in enumerate(rows, start=2):
        if all(value is None or str(value).strip() == "" for value in values):
            continue
        try:
            if len(values) != len(headers):
                raise ValueError("列数与表头不一致")
            record = normalize_row(dict(zip(headers, values)))
            key = (record.date, record.dish_id)
            if key in seen:
                raise ValueError("日期和菜品重复")
            if record.dish_id in names and names[record.dish_id] != record.dish_name:
                raise ValueError("同一菜品编号对应多个名称")
            seen.add(key)
            names[record.dish_id] = record.dish_name
            records.append(record)
        except (ValueError, ValidationError) as error:
            errors.append({"row": line, "message": str(error)})
    if errors:
        raise AppError(
            "数据校验失败，整批未写入", 422, {"error_count": len(errors), "rows": errors[:100]}
        )
    if not records:
        raise AppError("表格没有可导入的数据")
    return records


def _limited_rows(table, max_rows):
    for index, row in enumerate(table):
        if index >= max_rows:
            raise AppError(f"最多导入 {max_rows} 行，请分批处理", 413)
        yield list(row)


def normalize_row(row) -> SaleRecord:
    cleaned = {
        key: value for key, value in row.items() if value is not None and str(value).strip() != ""
    }
    for key in ("sales_quantity", "prepared_quantity", "waste_quantity"):
        if key in cleaned:
            try:
                number = Decimal(str(cleaned[key]))
                if (
                    not number.is_finite()
                    or number != number.to_integral_value()
                    or not 0 <= number <= 1000000
                ):
                    raise ValueError(f"{key} 必须是 0–1000000 的整数")
                cleaned[key] = int(number)
            except InvalidOperation as error:
                raise ValueError(f"{key} 必须是整数") from error
    for key in ("is_exam_week", "has_event"):
        if key in cleaned:
            text = str(cleaned[key]).strip().lower()
            if text not in {"true", "false", "1", "0", "yes", "no"}:
                raise ValueError(f"{key} 仅接受 true/false、1/0、yes/no")
            cleaned[key] = text in {"true", "1", "yes"}
    if isinstance(cleaned.get("date"), datetime):
        cleaned["date"] = cleaned["date"].date()
    elif not isinstance(cleaned.get("date"), date):
        cleaned["date"] = date.fromisoformat(str(cleaned.get("date", "")))
    for key in ("dish_id", "dish_name", "weather"):
        if key in cleaned:
            cleaned[key] = str(cleaned[key]).strip()
    return SaleRecord.model_validate(cleaned)
